import io
from unittest.mock import patch

from django.test import SimpleTestCase, TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from examc_app.models import PrepStudent
from examc_app.services.oasis import OasisError, get_course_students_scipers
from examc_app.services.person_directory import DirectoryPerson, get_people_by_sciper
from examc_app.services.student.prep_import import (
    StudentsFileError,
    build_students_export,
    build_students_template,
    correct_prep_student,
    load_students_file,
    load_students_from_oasis,
    read_students_file,
    replace_prep_students,
)
from examc_app.tests.helpers.models import create_mock_exam, create_mock_user

HEADER = ["SCIPER", "ROOM", "SEAT"]

DIRECTORY = {
    "123456": DirectoryPerson("123456", "Ada", "Lovelace", "ada.lovelace@epfl.ch", "MX"),
    "234567": DirectoryPerson("234567", "Alan", "Turing", None),
    "345678": DirectoryPerson("345678", "Émilie", "du Châtelet", None, "PH"),
    "456789": DirectoryPerson("456789", "Charles", "Babbage", None, "IN"),
}


def make_xlsx(rows: list[list]) -> bytes:
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def fake_directory(scipers):
    return {sciper: DIRECTORY[sciper] for sciper in scipers if sciper in DIRECTORY}


class ReadStudentsFileTestCase(SimpleTestCase):
    def test_xlsx_numbers_are_read_as_integers_text(self):
        content = make_xlsx([
            HEADER,
            [234567.0, None, 12],
            [None, None, None],
            [123456, "CO 1", "A1"],
        ])

        students = read_students_file("students.xlsx", content)

        # Copy numbers in the order of the rows, empty rows ignored
        self.assertEqual([(s["copy_no"], s["sciper"], s["row"]) for s in students], [(1, "234567", 2), (2, "123456", 4)])
        self.assertEqual(students[0]["seat"], "12")
        self.assertEqual(students[0]["room"], "")

    def test_headers_ignore_case_spaces_asterisks_and_order(self):
        content = make_xlsx([["seat *", "Sciper*"], ["A1", 123456]])

        students = read_students_file("students.xlsx", content)

        self.assertEqual(students[0]["sciper"], "123456")
        self.assertEqual(students[0]["seat"], "A1")

    def test_all_row_errors_are_reported(self):
        content = make_xlsx([
            HEADER,
            [123456, None, "A1"],
            [12345, None, None],
            [123456, None, "A3"],
            [None, "CO 1", "A4"],
        ])

        with self.assertRaises(StudentsFileError) as context:
            read_students_file("students.xlsx", content)

        errors = context.exception.errors
        self.assertEqual(len(errors), 3)
        self.assertIn("Row 3", errors[0])
        self.assertNotIn("seat", errors[0])
        self.assertIn("SCIPER '12345' must have 6 digits", errors[0])
        self.assertIn("SCIPER 123456 already used in row 2", errors[1])
        self.assertEqual(errors[2], "Row 5: sciper is empty.")

    def test_missing_required_columns(self):
        content = make_xlsx([["ID", "SCIPER", "NAME"], [1, 123456, "Ada Lovelace"]])

        with self.assertRaises(StudentsFileError) as context:
            read_students_file("students.xlsx", content)

        self.assertIn("Missing column(s) in row 1: SEAT", context.exception.errors[0])

    def test_too_long_value(self):
        content = make_xlsx([HEADER, [123456, "R" * 501, "A1"]])

        with self.assertRaises(StudentsFileError) as context:
            read_students_file("students.xlsx", content)

        self.assertIn("room is longer than 500 characters", context.exception.errors[0])

    def test_unsupported_or_broken_files(self):
        for name, content in [("students.csv", b"ID,SCIPER"), ("students.xls", b"x"), ("students.xlsx", b"not a zip")]:
            with self.subTest(name=name), self.assertRaises(StudentsFileError):
                read_students_file(name, content)

    def test_template_headers_are_accepted(self):
        template = load_workbook(io.BytesIO(build_students_template()))
        self.assertEqual(template.sheetnames, ["Students", "Help"])
        headers = [cell.value for cell in template["Students"][1]]
        self.assertEqual(headers, ["SCIPER *", "ROOM", "SEAT *"])
        self.assertEqual(template["Students"]["A1"].font.name, "Arial")

        # Empty template: only the header row
        with self.assertRaises(StudentsFileError) as context:
            read_students_file("students_template.xlsx", build_students_template())
        self.assertIn("no students", context.exception.errors[0])

        # Filled template
        template["Students"].append([123456, None, "A1"])
        buffer = io.BytesIO()
        template.save(buffer)
        self.assertEqual(len(read_students_file("students_template.xlsx", buffer.getvalue())), 1)


@patch("examc_app.services.student.prep_import.get_people_by_sciper", side_effect=fake_directory)
class LoadStudentsFileTestCase(SimpleTestCase):
    def test_names_and_email_come_from_the_directory(self, _directory):
        students, warnings = load_students_file("students.xlsx", make_xlsx([
            HEADER, [123456, "CO 1", "A1"], [234567, None, "A2"],
        ]))

        self.assertEqual(warnings, [])
        self.assertEqual(
            [(s["first_name"], s["last_name"], s["email"], s["section"]) for s in students],
            [("Ada", "Lovelace", "ada.lovelace@epfl.ch", "MX"), ("Alan", "Turing", None, None)],
        )

    def test_unknown_sciper_is_imported_to_correct(self, _directory):
        students, warnings = load_students_file("students.xlsx", make_xlsx([
            HEADER, [999999, None, "A1"], [123456, None, "A2"],
        ]))

        # Kept at its place in the file
        self.assertEqual([(s["copy_no"], s["sciper"], s["last_name"], s.get("needs_correction", False))
                          for s in students], [(1, "999999", "", True), (2, "123456", "Lovelace", False)])
        self.assertEqual(warnings, ["Row 2: SCIPER 999999 not found in the EPFL directory, correct it in the table."])

    def test_rows_without_seat_are_imported_with_one_warning(self, _directory):
        students, warnings = load_students_file("students.xlsx", make_xlsx([
            HEADER, [123456, "CO 1", None], [234567, None, "A2"], [345678, None, ""],
        ]))

        self.assertEqual([s["seat"] for s in students], ["", "A2", ""])
        self.assertEqual(warnings, ["2 student(s) without seat (row 2, 4): give them one in the table (in yellow)."])

    def test_unknown_sciper_with_names_in_the_file_is_kept(self, _directory):
        students, warnings = load_students_file("students.xlsx", make_xlsx([
            HEADER + ["LAST NAME", "FIRST NAME", "EMAIL"],
            [999999, None, "A1", "Hopper", "Grace", "grace@example.com"],
            [123456, None, "A2", "Wrong", "Name", None],
        ]))

        self.assertEqual(warnings, [])
        self.assertEqual(
            [(s["last_name"], s["first_name"], s["email"], s.get("needs_correction", False)) for s in students],
            # The directory wins when it knows the SCIPER
            [("Hopper", "Grace", "grace@example.com", False), ("Lovelace", "Ada", "ada.lovelace@epfl.ch", False)],
        )


class FakeExam:
    code = "AR-101"

    class year:
        code = "2026-2027"


@patch("examc_app.services.student.prep_import.get_people_by_sciper", side_effect=fake_directory)
class LoadStudentsFromOasisTestCase(SimpleTestCase):
    @patch("examc_app.services.student.prep_import.get_course_students_scipers",
           return_value=["123456", "999999", "345678", "234567", "456789"])
    def test_sorted_by_last_name_and_numbered(self, oasis, _directory):
        students, warnings = load_students_from_oasis(FakeExam)

        oasis.assert_called_once_with("2026-2027", "AR-101")
        # "du Châtelet" between Babbage and Lovelace: case and accents ignored; the SCIPER not found comes last
        self.assertEqual([(s["copy_no"], s["last_name"], s.get("needs_correction", False)) for s in students],
                         [(1, "Babbage", False), (2, "du Châtelet", False), (3, "Lovelace", False), (4, "Turing", False),
                          (5, "", True)])
        self.assertEqual((students[0]["seat"], students[0]["section"]), ("", "IN"))
        self.assertEqual(warnings, ["SCIPER 999999 not found in the EPFL directory, correct it in the table."])

    @patch("examc_app.services.student.prep_import.get_course_students_scipers", return_value=[])
    def test_no_enrolled_student(self, _oasis, _directory):
        with self.assertRaises(StudentsFileError) as context:
            load_students_from_oasis(FakeExam)

        self.assertIn("No student is enrolled in AR-101 for 2026-2027", context.exception.errors[0])


class GetCourseStudentsScipersTestCase(SimpleTestCase):
    @patch("examc_app.services.oasis._get_list")
    def test_deduplicated_in_order(self, get_list):
        row = {"anneeAcademique": "2026-2027", "codeCours": "AR-101", "titreCoursFr": "Cours"}
        get_list.return_value = [{**row, "sciper": "234567"}, {**row, "sciper": 123456}, {**row, "sciper": "234567"}]

        self.assertEqual(get_course_students_scipers("2026-2027", "AR-101"), ["234567", "123456"])
        get_list.assert_called_once_with("inscription-cours",
                                         params={"annee-academique": "2026-2027", "code-cours": "AR-101"})

    @patch("examc_app.services.oasis._get_list")
    def test_other_course_or_invalid_sciper_is_refused(self, get_list):
        for row in ({"anneeAcademique": "2026-2027", "codeCours": "AR-102", "sciper": "123456"},
                    {"anneeAcademique": "2026-2027", "codeCours": "AR-101", "sciper": "12345"}):
            get_list.return_value = [row]
            with self.subTest(row=row), self.assertRaises(OasisError):
                get_course_students_scipers("2026-2027", "AR-101")


class GetPeopleBySciperTestCase(SimpleTestCase):
    @patch("examc_app.services.person_directory.LDAP_search")
    def test_one_query_per_batch_and_first_entry_kept(self, ldap_search):
        ldap_search.return_value = [
            {"dn": "cn=Ada Lovelace,ou=lab,ou=sti,o=epfl,c=ch",
             "attributes": {"uniqueIdentifier": ["123456"], "givenName": ["Ada"], "sn": ["Lovelace"],
                            "mail": ["ada.lovelace@epfl.ch"], "ou": ["LAB-X"]}},
            # Second accreditation of the same person: her student entry gives the section
            {"dn": "cn=Ada Lovelace,ou=mx-ba1,ou=mx-s,ou=etu,o=epfl,c=ch",
             "attributes": {"uniqueIdentifier": ["123456"], "givenName": ["Other"], "sn": ["Entry"], "mail": [],
                            "ou": ["MX-BA1", "Section Science et génie des matériaux - Bachelor semestre 1"]}},
            {"dn": "cn=Alan Turing,ou=cepro,ou=avp,o=epfl,c=ch",
             "attributes": {"uniqueIdentifier": ["234567"], "givenName": ["Alan"], "sn": ["Turing"],
                            "mail": ["not-an-email"], "ou": ["AVP-E-CEPRO"]}},
        ]

        people = get_people_by_sciper(["234567", 123456, "123456"])

        ldap_search.assert_called_once_with(pattern_search="(|(uniqueIdentifier=123456)(uniqueIdentifier=234567))")
        self.assertEqual(people["123456"], DirectoryPerson("123456", "Ada", "Lovelace", "ada.lovelace@epfl.ch", "MX"))
        self.assertIsNone(people["234567"].email)
        # Not a student: no section
        self.assertIsNone(people["234567"].section)

    def test_invalid_sciper_is_refused_before_ldap(self):
        with self.assertRaises(ValueError):
            get_people_by_sciper(["12345)(uid=*"])


class ReplacePrepStudentsTestCase(TestCase):
    @patch("examc_app.services.student.prep_import.get_people_by_sciper", side_effect=fake_directory)
    def test_replaces_the_exam_students_only(self, _directory):
        exam = create_mock_exam()
        other_exam = create_mock_exam(code="OTHER")
        PrepStudent.objects.create(exam=exam, copy_no=9, sciper=999999, first_name="Old", last_name="Student", seat="Z")
        PrepStudent.objects.create(exam=other_exam, copy_no=1, sciper=111111, first_name="Keep", last_name="Me", seat="B")
        students, _ = load_students_file("students.xlsx", make_xlsx([
            HEADER, [123456, "CO 1", "A1"], [234567, None, "A2"],
        ]))

        result = replace_prep_students(exam, students)

        self.assertEqual((result.imported, result.replaced), (2, 1))
        self.assertEqual(
            list(exam.prepStudents.order_by("copy_no").values_list(
                "copy_no", "sciper", "last_name", "email", "section", "room")),
            [(1, 123456, "Lovelace", "ada.lovelace@epfl.ch", "MX", "CO 1"), (2, 234567, "Turing", None, None, None)],
        )
        self.assertEqual(other_exam.prepStudents.count(), 1)
        self.assertEqual(PrepStudent.history.filter(exam_id=exam.pk, history_type="+").count(), 3)


class StudentsExportTestCase(TestCase):
    def setUp(self):
        self.exam = create_mock_exam()
        PrepStudent.objects.create(exam=self.exam, copy_no=2, sciper=123456, first_name="Ada", last_name="Lovelace",
                                   email="ada.lovelace@epfl.ch", section="MX", room="CO 1", seat="A2")
        PrepStudent.objects.create(exam=self.exam, copy_no=1, sciper=999999, first_name="Grace", last_name="Hopper",
                                   seat="A1")
        PrepStudent.objects.create(exam=create_mock_exam(code="OTHER"), copy_no=1, sciper=111111, first_name="F",
                                   last_name="Other", seat="B")

    def test_rows_in_copy_order(self):
        sheet = load_workbook(io.BytesIO(build_students_export(self.exam)))["Students"]

        self.assertEqual(list(sheet.iter_rows(values_only=True)), [
            ("SCIPER *", "ROOM", "SEAT *", "LAST NAME", "FIRST NAME", "EMAIL", "SECTION"),
            (999999, None, "A1", "Hopper", "Grace", None, None),
            (123456, "CO 1", "A2", "Lovelace", "Ada", "ada.lovelace@epfl.ch", "MX"),
        ])

    @patch("examc_app.services.student.prep_import.get_people_by_sciper", side_effect=fake_directory)
    def test_imported_back_unchanged(self, _directory):
        before = list(self.exam.prepStudents.order_by("copy_no").values_list(
            "copy_no", "sciper", "last_name", "first_name", "email", "section", "room", "seat", "needs_correction"))

        students, warnings = load_students_file("students.xlsx", build_students_export(self.exam))
        replace_prep_students(self.exam, students, warnings)

        self.assertEqual(warnings, [])
        self.assertEqual(list(self.exam.prepStudents.order_by("copy_no").values_list(
            "copy_no", "sciper", "last_name", "first_name", "email", "section", "room", "seat", "needs_correction")),
            before)

    def test_download(self):
        user = create_mock_user()
        user.is_superuser = True
        user.save()
        self.client.force_login(user)

        response = self.client.get(reverse("download_prep_students_export", kwargs={"exam_pk": self.exam.pk}))

        self.assertEqual(response.status_code, 200)
        self.assertIn(f'filename="students_{self.exam.code}_', response["Content-Disposition"])
        self.assertEqual(len(read_students_file("students.xlsx", response.content)), 2)


class PrepStudentApiTestCase(TestCase):
    def test_lists_all_the_exam_students_by_copy_number(self):
        exam = create_mock_exam()
        for copy_no in (3, 1, 2):
            PrepStudent.objects.create(exam=exam, copy_no=copy_no, sciper=100000 + copy_no, first_name="F",
                                       last_name=f"L{copy_no}", seat="")
        PrepStudent.objects.create(exam=create_mock_exam(code="OTHER"), copy_no=9, sciper=999999, first_name="F",
                                   last_name="Other", seat="")
        user = create_mock_user()
        user.is_superuser = True
        user.save()
        self.client.force_login(user)

        response = self.client.get(reverse("api-exam-prep-students-list", kwargs={"exam_pk": exam.pk}))

        self.assertEqual(response.status_code, 200)
        self.assertEqual([row["copy_no"] for row in response.json()["data"]], [1, 2, 3])


@patch("examc_app.services.student.prep_import.get_people_by_sciper", side_effect=fake_directory)
class CorrectPrepStudentTestCase(TestCase):
    def setUp(self):
        self.exam = create_mock_exam()
        self.student = PrepStudent.objects.create(exam=self.exam, copy_no=1, sciper=999999, first_name="",
                                                  last_name="", seat="A1", needs_correction=True)

    def test_new_sciper_found_fills_from_the_directory(self, _directory):
        correct_prep_student(self.student, {"sciper": 123456, "last_name": "typed", "room": "CO 1"})

        self.student.refresh_from_db()
        self.assertEqual((self.student.sciper, self.student.last_name, self.student.section, self.student.room,
                          self.student.needs_correction), (123456, "Lovelace", "MX", "CO 1", False))

    def test_person_missing_from_the_directory_added_manually(self, _directory):
        correct_prep_student(self.student, {"sciper": 888888, "last_name": "Hopper", "first_name": "Grace"})

        self.student.refresh_from_db()
        self.assertEqual((self.student.sciper, self.student.last_name, self.student.needs_correction),
                         (888888, "Hopper", False))

    def test_not_found_without_names_is_refused(self, _directory):
        with self.assertRaises(StudentsFileError):
            correct_prep_student(self.student, {"sciper": 888888})

        self.student.refresh_from_db()
        self.assertEqual((self.student.sciper, self.student.needs_correction), (999999, True))


class PrepStudentPatchApiTestCase(TestCase):
    def setUp(self):
        self.exam = create_mock_exam()
        self.student = PrepStudent.objects.create(exam=self.exam, copy_no=1, sciper=999999, first_name="",
                                                  last_name="", seat="A1", needs_correction=True)
        PrepStudent.objects.create(exam=self.exam, copy_no=2, sciper=234567, first_name="Alan", last_name="Turing",
                                   seat="A2")
        user = create_mock_user()
        user.is_superuser = True
        user.save()
        self.client.force_login(user)
        self.url = reverse("api-exam-prep-students-detail", kwargs={"exam_pk": self.exam.pk, "pk": self.student.pk})

    @patch("examc_app.services.student.prep_import.get_people_by_sciper", side_effect=fake_directory)
    def test_correct_the_sciper(self, _directory):
        response = self.client.patch(self.url, {"sciper": "123456", "email": "", "room": ""},
                                     content_type="application/json")

        self.assertEqual(response.status_code, 200)
        self.assertEqual((response.json()["last_name"], response.json()["needs_correction"]), ("Lovelace", False))

    def test_sciper_already_used_or_invalid(self):
        for sciper, message in (("234567", "already used"), ("12345", "6 digits")):
            response = self.client.patch(self.url, {"sciper": sciper}, content_type="application/json")
            with self.subTest(sciper=sciper):
                self.assertEqual(response.status_code, 400)
                self.assertIn(message, response.json()["sciper"][0])

    def test_finalized_exam_is_locked(self):
        self.exam.is_finalized = True
        self.exam.save()

        response = self.client.patch(self.url, {"seat": "B1"}, content_type="application/json")

        self.assertEqual(response.status_code, 403)

    def test_final_generation_blocked_while_students_to_correct(self):
        response = self.client.get(reverse("generate_final_exam_files_start", kwargs={"exam_pk": self.exam.pk}))

        self.assertEqual(response.status_code, 400)
        self.assertIn("1 student(s) of the Students page must be corrected", response.json()["error"])

    def test_final_generation_blocked_while_students_without_seat(self):
        self.student.needs_correction = False
        self.student.seat = ""
        self.student.save()

        response = self.client.get(reverse("generate_final_exam_files_start", kwargs={"exam_pk": self.exam.pk}))

        self.assertEqual(response.status_code, 400)
        self.assertIn("1 student(s) of the Students page have no seat", response.json()["error"])

