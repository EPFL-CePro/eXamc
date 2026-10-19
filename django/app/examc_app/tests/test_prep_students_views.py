import io
from unittest.mock import patch

from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook

from examc_app.models import PrepStudent
from examc_app.services.oasis import OasisError
from examc_app.services.person_directory import DirectoryPerson, PersonDirectoryError
from examc_app.services.student.prep_import import StudentsFileError
from examc_app.tests.helpers.prep import PrepViewsTestCase

DIRECTORY = {
    "123456": DirectoryPerson("123456", "Ada", "Lovelace", "ada.lovelace@epfl.ch", "MX"),
    "234567": DirectoryPerson("234567", "Alan", "Turing", None),
}


def fake_directory(scipers):
    return {sciper: DIRECTORY[sciper] for sciper in scipers if sciper in DIRECTORY}


def xlsx(rows) -> bytes:
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class StudentsViewsTestCase(PrepViewsTestCase):
    def student(self, copy_no, sciper, last_name="Doe"):
        return PrepStudent.objects.create(exam=self.exam, copy_no=copy_no, sciper=sciper, first_name="John",
                                          last_name=last_name, seat=f"S{copy_no}")

    def names(self):
        return list(self.exam.prepStudents.order_by("copy_no").values_list("copy_no", "sciper", "last_name"))


class StudentsTemplateTestCase(StudentsViewsTestCase):
    def test_template_is_an_excel_file_with_the_columns_to_fill(self):
        response = self.client.get(self.url("download_prep_students_template"))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"],
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn("students_template.xlsx", response["Content-Disposition"])
        sheet = load_workbook(io.BytesIO(response.content))["Students"]
        header = [cell.value for cell in sheet[1]]
        self.assertIn("SCIPER *", header)
        self.assertIn("SEAT *", header)


@patch("examc_app.services.student.prep_import.get_people_by_sciper", side_effect=fake_directory)
class ImportStudentsXlsxTestCase(StudentsViewsTestCase):
    def upload(self, content=None, name="students.xlsx"):
        data = {} if content is None else {"students_file": SimpleUploadedFile(name, content)}
        return self.client.post(self.url("import_prep_students_xlsx"), data)

    def test_replaces_the_students_by_those_of_the_file(self, _directory):
        self.student(1, 999999, "Old")

        response = self.upload(xlsx([["SCIPER", "ROOM", "SEAT"], [123456, "CO 1", "A1"], [234567, "CO 1", "A2"]]))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), {"imported": 2, "replaced": 1, "warnings": []})
        self.assertEqual(self.names(), [(1, 123456, "Lovelace"), (2, 234567, "Turing")])
        self.assertEqual(self.exam.prepStudents.get(copy_no=1).room, "CO 1")

    def test_warns_about_the_sciper_not_found_in_the_directory(self, _directory):
        response = self.upload(xlsx([["SCIPER", "ROOM", "SEAT"], [999999, "CO 1", "A1"]]))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["imported"], 1)
        self.assertIn("SCIPER 999999 not found in the EPFL directory", response.json()["warnings"][0])
        self.assertTrue(self.exam.prepStudents.get().needs_correction)

    def test_a_file_with_problems_changes_nothing(self, _directory):
        self.student(1, 999999, "Old")

        response = self.upload(xlsx([["SCIPER", "ROOM", "SEAT"], ["not a sciper", "CO 1", "A1"]]))

        self.assertEqual(response.status_code, 400)
        self.assertTrue(response.json()["errors"])
        self.assertEqual(self.names(), [(1, 999999, "Old")])

    def test_a_file_that_is_not_an_excel_file_is_refused(self, _directory):
        response = self.upload(b"SCIPER;ROOM;SEAT\n123456;CO 1;A1\n", name="students.csv")

        self.assertEqual(response.status_code, 400)
        self.assertIn("is not an Excel (.xlsx) file", response.json()["errors"][0])

    def test_without_a_file(self, _directory):
        response = self.upload()

        self.assertEqual((response.status_code, response.json()), (400, {"errors": ["No file selected."]}))

    def test_directory_that_cannot_be_reached(self, directory):
        directory.side_effect = PersonDirectoryError("down")
        self.student(1, 999999, "Old")

        with self.assertLogs("examc_app.views.preparation.students", level="ERROR"):
            response = self.upload(xlsx([["SCIPER", "ROOM", "SEAT"], [123456, "CO 1", "A1"]]))

        self.assertEqual(response.status_code, 503)
        self.assertIn("EPFL directory could not be reached", response.json()["errors"][0])
        self.assertEqual(self.names(), [(1, 999999, "Old")])

    def test_get_is_refused(self, _directory):
        self.assertEqual(self.client.get(self.url("import_prep_students_xlsx")).status_code, 400)

    def test_locked_exam_is_refused(self, _directory):
        self.student(1, 999999, "Old")
        self.lock()

        response = self.upload(xlsx([["SCIPER", "ROOM", "SEAT"], [123456, "CO 1", "A1"]]))

        self.assertEqual(response.status_code, 403)
        self.assertEqual(self.names(), [(1, 999999, "Old")])


class ImportStudentsApiTestCase(StudentsViewsTestCase):
    LOADED = [{"copy_no": 1, "sciper": "123456", "first_name": "Ada", "last_name": "Lovelace", "email": "",
               "section": "MX", "room": "", "seat": ""}]

    def run_import(self, **patches):
        with patch("examc_app.views.preparation.students.load_students_from_oasis", **patches) as load:
            response = self.client.post(self.url("import_prep_students_api"))
        return response, load

    def test_replaces_the_students_by_those_enrolled_in_the_course(self):
        self.student(1, 999999, "Old")

        response, load = self.run_import(return_value=(self.LOADED, ["a warning"]))

        load.assert_called_once()
        self.assertEqual(response.json(), {"imported": 1, "replaced": 1, "warnings": ["a warning"]})
        self.assertEqual(self.names(), [(1, 123456, "Lovelace")])

    def test_problems_of_the_enrolment_are_listed(self):
        self.student(1, 999999, "Old")

        response, _ = self.run_import(side_effect=StudentsFileError(["No student is enrolled in this course."]))

        self.assertEqual((response.status_code, response.json()), (400, {"errors": ["No student is enrolled in this course."]}))
        self.assertEqual(self.names(), [(1, 999999, "Old")])

    def test_oasis_or_the_directory_that_cannot_be_reached(self):
        self.student(1, 999999, "Old")

        for error, name in ((OasisError("down"), "IS-Academia"), (PersonDirectoryError("down"), "EPFL directory")):
            with self.subTest(error=name):
                with self.assertLogs("examc_app.views.preparation.students", level="ERROR"):
                    response, _ = self.run_import(side_effect=error)

                self.assertEqual(response.status_code, 503)
                self.assertIn(name, response.json()["errors"][0])
        self.assertEqual(self.names(), [(1, 999999, "Old")])

    def test_get_is_refused(self):
        self.assertEqual(self.client.get(self.url("import_prep_students_api")).status_code, 400)

    def test_locked_exam_is_refused(self):
        self.student(1, 999999, "Old")
        self.lock()

        response, load = self.run_import(return_value=(self.LOADED, []))

        self.assertEqual(response.status_code, 403)
        load.assert_not_called()
        self.assertEqual(self.names(), [(1, 999999, "Old")])
