import io
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font
from openpyxl.utils.indexed_list import IndexedList
from openpyxl.worksheet.datavalidation import DataValidation
from simple_history.utils import bulk_create_with_history

from examc_app.models import Exam, PrepStudent
from examc_app.services.oasis import get_course_students_scipers
from examc_app.services.person_directory import get_people_by_sciper


# Column header shown to the users (template file order) -> PrepStudent field.
# Names, email and section come from the EPFL directory (see complete_from_directory),
# the copy numbers follow the order of the rows.
HEADER_FIELDS = {
    "SCIPER": "sciper",
    "ROOM": "room",
    "SEAT": "seat",
}
# Required columns. A row without seat is still imported, the seat is given later in the table
# (the final files cannot be generated before, see generate_final_exam_files_start).
REQUIRED_FIELDS = ("sciper", "seat")
ROW_REQUIRED_FIELDS = ("sciper",)
# Header -> (content, example), shown in the template file
HEADER_HELP = {
    "SCIPER": ("SCIPER: 6 digits, unique in the exam", 123456),
    "ROOM": ("Room", "CO 1"),
    "SEAT": ("Seat; a row without seat is imported, the seat is to give in eXamc before the generation", "A12"),
    "LAST NAME": ("Last name, from the EPFL directory", "Lovelace"),
    "FIRST NAME": ("First name, from the EPFL directory", "Ada"),
    "EMAIL": ("Email, from the EPFL directory", "ada.lovelace@epfl.ch"),
    "SECTION": ("Section, from the EPFL directory", "MX"),
}
# Columns of the exported list (build_students_export) after HEADER_FIELDS. When read back, they are only used
# for a SCIPER missing from the EPFL directory (a person added manually, see correct_prep_student).
INFO_HEADER_FIELDS = {
    "LAST NAME": "last_name",
    "FIRST NAME": "first_name",
    "EMAIL": "email",
    "SECTION": "section",
}
TEXT_FIELDS = ("room", "seat", *INFO_HEADER_FIELDS.values())

SCIPER_RE = re.compile(r"^\d{6}$")


class StudentsFileError(Exception):
    """The file cannot be imported: `errors` lists the problems, for the user."""

    def __init__(self, errors: list[str]):
        super().__init__("; ".join(errors))
        self.errors = errors


@dataclass(frozen=True)
class StudentsImportResult:
    imported: int
    replaced: int
    # Problems that did not prevent the import, for the user
    warnings: tuple[str, ...] = ()


def _normalize_header(value) -> str:
    """'Last name', 'LAST_NAME', ' lastname ' -> 'lastname'"""
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z]", "", text.lower())


def _cell_text(value) -> str:
    """Spreadsheet cell -> text: Excel stores 123456 as 123456.0"""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _read_xlsx_rows(content: bytes) -> list[list]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise StudentsFileError(["The file is not a valid Excel (.xlsx) file."])

    try:
        return [list(row) for row in workbook.worksheets[0].iter_rows(values_only=True)]
    finally:
        workbook.close()


def read_students_file(file_name: str, content: bytes) -> list[dict]:
    """
    Reads the students of the first sheet of an .xlsx file, whose first row holds the column headers
    (HEADER_FIELDS, case, spaces and underscores ignored). The copy numbers are 1, 2, 3... in the order of the rows.
    Raises StudentsFileError with all the problems found.
    """
    if Path(file_name).suffix.lower() != ".xlsx":
        raise StudentsFileError([f"'{file_name}' is not an Excel (.xlsx) file: "
                                 "download the template, fill it and import it."])

    rows = _read_xlsx_rows(content)
    # Rows left empty (often at the end of a spreadsheet) are ignored
    rows = [(number, row) for number, row in enumerate(rows, start=1) if any(_cell_text(v) for v in row)]
    if not rows:
        raise StudentsFileError(["The file is empty."])

    header_number, header = rows[0]
    fields_by_header = {_normalize_header(h): field for h, field in {**HEADER_FIELDS, **INFO_HEADER_FIELDS}.items()}
    field_by_index = {}
    for index, value in enumerate(header):
        field = fields_by_header.get(_normalize_header(value))
        if field and field not in field_by_index.values():
            field_by_index[index] = field

    missing = [h for h, field in HEADER_FIELDS.items()
               if field in REQUIRED_FIELDS and field not in field_by_index.values()]
    if missing:
        raise StudentsFileError([f"Missing column(s) in row {header_number}: {', '.join(missing)}. "
                                 f"Expected columns: {', '.join(HEADER_FIELDS)}."])

    max_lengths = {field: PrepStudent._meta.get_field(field).max_length for field in TEXT_FIELDS}

    errors = []
    students = []
    seen_sciper = {}

    for number, row in rows[1:]:
        student = {field: _cell_text(row[index]) if index < len(row) else "" for index, field in field_by_index.items()}
        row_errors = [f"{field.replace('_', ' ')} is empty" for field in ROW_REQUIRED_FIELDS if not student.get(field)]

        if student.get("sciper"):
            if not SCIPER_RE.match(student["sciper"]):
                row_errors.append(f"SCIPER '{student['sciper']}' must have 6 digits")
            elif student["sciper"] in seen_sciper:
                row_errors.append(f"SCIPER {student['sciper']} already used in row {seen_sciper[student['sciper']]}")
            else:
                seen_sciper[student["sciper"]] = number

        for field, max_length in max_lengths.items():
            if len(student.get(field, "")) > max_length:
                row_errors.append(f"{field.replace('_', ' ')} is longer than {max_length} characters")

        if row_errors:
            errors.append(f"Row {number}: {', '.join(row_errors)}.")
        else:
            student["row"] = number
            student["copy_no"] = len(students) + 1
            students.append(student)

    if errors:
        raise StudentsFileError(errors)
    if not students:
        raise StudentsFileError(["The file has a header row but no students."])
    return students


def complete_from_directory(students: list[dict]) -> list[dict]:
    """
    Adds first name, last name, email and section from the EPFL directory to the students that are found in it,
    and returns the students that are not. Raises PersonDirectoryError if the directory cannot be reached.
    """
    people = get_people_by_sciper(student["sciper"] for student in students)

    for student in students:
        person = people.get(student["sciper"])
        if person:
            student.update(first_name=person.first_name, last_name=person.last_name, email=person.email,
                           section=person.section)
    return [student for student in students if student["sciper"] not in people]


def _mark_to_correct(students: list[dict]) -> None:
    """Students whose SCIPER is not in the EPFL directory: imported without names, to be corrected in the table."""
    for student in students:
        student.update(first_name="", last_name="", email=None, section=None, needs_correction=True)


def _has_names(student: dict) -> bool:
    return bool(student.get("last_name") and student.get("first_name"))


def load_students_file(file_name: str, content: bytes) -> tuple[list[dict], list[str]]:
    """
    Students of an imported .xlsx file, completed from the EPFL directory, ready for replace_prep_students;
    and the warnings for the user. A SCIPER not found in the directory is imported to be corrected (see
    _mark_to_correct), unless the file gives its last and first names (a person added manually, in an exported
    list): these names are kept. A row without seat is imported, the seat is to give in the table.
    """
    students = read_students_file(file_name, content)
    missing = [student for student in complete_from_directory(students) if not _has_names(student)]
    _mark_to_correct(missing)
    warnings = [f"Row {student['row']}: SCIPER {student['sciper']} not found in the EPFL directory, "
                "correct it in the table." for student in missing]
    without_seat = [str(student["row"]) for student in students if not student.get("seat")]
    if without_seat:
        shown = ", ".join(without_seat[:20]) + (", ..." if len(without_seat) > 20 else "")
        warnings.append(f"{len(without_seat)} student(s) without seat (row {shown}): "
                        "give them one in the table (in yellow).")
    return students, warnings


def _name_sort_key(student: dict) -> tuple[str, str]:
    """Last name then first name, ignoring case and accents (É sorted with E)."""
    def key(text):
        return unicodedata.normalize("NFKD", text or "").encode("ascii", "ignore").decode().casefold()
    return key(student["last_name"]), key(student["first_name"])


def load_students_from_oasis(exam: Exam) -> tuple[list[dict], list[str]]:
    """
    Students enrolled in the exam course (IS-Academia, through OASIS), completed from the EPFL directory and
    sorted by last name, with copy numbers 1, 2, 3... in that order; and the warnings for the user.
    The SCIPERs not found in the directory come last, to be corrected (see _mark_to_correct).
    Raises StudentsFileError when no student is enrolled, OasisError or PersonDirectoryError when a
    service cannot be reached.
    """
    scipers = get_course_students_scipers(exam.year.code, exam.code)
    if not scipers:
        raise StudentsFileError([f"No student is enrolled in {exam.code} for {exam.year.code} in IS-Academia."])

    # No room nor seat in IS-Academia
    students = [{"sciper": sciper, "room": "", "seat": ""} for sciper in scipers]
    missing = complete_from_directory(students)
    warnings = [f"SCIPER {student['sciper']} not found in the EPFL directory, correct it in the table."
                for student in missing]

    found = sorted((student for student in students if student not in missing), key=_name_sort_key)
    _mark_to_correct(missing)
    students = found + missing
    for copy_no, student in enumerate(students, start=1):
        student["copy_no"] = copy_no
    return students, warnings


def replace_prep_students(exam: Exam, students: list[dict], warnings: list[str] = ()) -> StudentsImportResult:
    """
    Replaces all the prep students of the exam by `students` (from load_students_file or load_students_from_oasis),
    all or nothing.
    """
    with transaction.atomic():
        current = PrepStudent.objects.filter(exam=exam)
        replaced = current.count()
        current.delete()
        bulk_create_with_history([
            PrepStudent(
                exam=exam,
                copy_no=int(student["copy_no"]),
                sciper=int(student["sciper"]),
                last_name=student["last_name"],
                first_name=student["first_name"],
                email=student.get("email") or None,
                section=student.get("section") or None,
                room=student.get("room") or None,
                seat=student.get("seat") or "",
                needs_correction=student.get("needs_correction", False),
            )
            for student in students
        ], PrepStudent)
    return StudentsImportResult(imported=len(students), replaced=replaced, warnings=tuple(warnings))


def _students_workbook(headers: dict[str, str]) -> Workbook:
    """
    Workbook for read_students_file: a "Students" sheet whose first row holds `headers` (required ones marked
    with "*", a comment on each) and a "Help" sheet describing the columns. Plain Arial, no styling.
    """
    workbook = Workbook()
    # Default font of the cells (openpyxl has no public API for it): Arial instead of Calibri
    arial = Font(name="Arial", size=11, family=2)
    workbook._fonts = IndexedList([arial])
    workbook._named_styles["Normal"].font = arial

    sheet = workbook.active
    sheet.title = "Students"

    for column, (header, field) in enumerate(headers.items(), start=1):
        required = field in REQUIRED_FIELDS
        content, _ = HEADER_HELP[header]
        # The "*" is ignored when reading the headers (see _normalize_header)
        cell = sheet.cell(row=1, column=column, value=f"{header} *" if required else header)
        cell.comment = Comment(f"{'Required' if required else 'Optional'}. {content}.", "eXamc")
        sheet.column_dimensions[cell.column_letter].width = 16
    sheet.freeze_panes = "A2"

    # Excel warns on a wrong SCIPER while typing
    sciper_validation = DataValidation(type="whole", operator="between", formula1="100000", formula2="999999",
                                       showErrorMessage=True, errorTitle="SCIPER", error="The SCIPER must have 6 digits.")
    sheet.add_data_validation(sciper_validation)
    sciper_validation.add("A2:A5000")

    help_sheet = workbook.create_sheet("Help")
    help_sheet.append(["Only the first sheet (Students) is imported. Keep its first row unchanged, "
                       "one student per row. Importing replaces all the students of the exam."])
    help_sheet.append(["The copy numbers are given in the order of the rows: 1 for the first student, 2 for the next..."])
    help_sheet.append(["* = required column. Names, emails and sections are taken from the EPFL directory with the SCIPER."])
    help_sheet.append([])
    help_sheet.append(["Column", "Required", "Content", "Example"])
    for header, field in headers.items():
        content, example = HEADER_HELP[header]
        help_sheet.append([header, "Required" if field in REQUIRED_FIELDS else "Optional", content, example])
    for letter, width in zip("ABCD", (14, 12, 48, 24)):
        help_sheet.column_dimensions[letter].width = width

    return workbook


def _workbook_bytes(workbook: Workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_students_template() -> bytes:
    """Empty Excel file to fill for read_students_file."""
    return _workbook_bytes(_students_workbook(HEADER_FIELDS))


def build_students_export(exam: Exam) -> bytes:
    """
    The students of the exam as an Excel file that read_students_file imports back, in the order of the copy
    numbers: to reorder them and fill rooms and seats in Excel. The names, email and section (INFO_HEADER_FIELDS)
    are for reading only, except for the people missing from the EPFL directory.
    """
    headers = {**HEADER_FIELDS, **INFO_HEADER_FIELDS}
    workbook = _students_workbook(headers)
    sheet = workbook["Students"]
    for student in PrepStudent.objects.filter(exam=exam).order_by("copy_no", "pk"):
        sheet.append([getattr(student, field) or None for field in headers.values()])
    for letter in "DEF":
        sheet.column_dimensions[letter].width = 28
    return _workbook_bytes(workbook)


def correct_prep_student(student: PrepStudent, data: dict) -> PrepStudent:
    """
    Saves the fields edited in the students table (`data`, validated by PrepStudentUpdateSerializer).
    A new SCIPER is looked up in the EPFL directory: when found, names, email and section come from it;
    otherwise the names typed by the user are kept (a person missing from the directory). The student needs
    no more correction once it has a last and a first name. Raises StudentsFileError when the SCIPER is not
    found and the names are missing, PersonDirectoryError if the directory cannot be reached.
    """
    sciper = str(data.get("sciper", student.sciper))
    if sciper != str(student.sciper):
        person = get_people_by_sciper([sciper]).get(sciper)
        if person:
            data = {**data, "first_name": person.first_name, "last_name": person.last_name,
                    "email": person.email, "section": person.section}

    for field, value in data.items():
        setattr(student, field, value)

    if not (student.last_name and student.first_name):
        raise StudentsFileError([f"SCIPER {sciper} is not in the EPFL directory: correct it, "
                                 "or fill in the last and first names to add this person manually."])

    student.needs_correction = False
    student.save()
    return student

